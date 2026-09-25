"""Gera planilhas de teste (.xlsx e .xls) com dados inventados."""

from __future__ import annotations

import io
from datetime import date, datetime

import xlwt
from openpyxl import Workbook

Linhas = list[list[object]]


def gerar_xlsx(linhas: Linhas, abas_extras: dict[str, Linhas] | None = None) -> bytes:
    livro = Workbook()
    planilha = livro.active
    planilha.title = "Extrato"
    for linha in linhas:
        planilha.append(linha)
    for nome, extras in (abas_extras or {}).items():
        aba = livro.create_sheet(nome)
        for linha in extras:
            aba.append(linha)
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def gerar_xls(linhas: Linhas) -> bytes:
    """Excel 97-2003, como o que alguns sistemas exportam."""
    livro = xlwt.Workbook()
    aba = livro.add_sheet("Extrato")
    formato_data = xlwt.easyxf(num_format_str="DD/MM/YYYY")
    for i, linha in enumerate(linhas):
        for j, valor in enumerate(linha):
            if isinstance(valor, (date, datetime)):
                aba.write(i, j, valor, formato_data)
            elif valor is not None:
                aba.write(i, j, valor)
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()
